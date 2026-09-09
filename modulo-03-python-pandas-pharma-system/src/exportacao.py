"""
exportacao.py
-------------
Exporta relatorio Excel multi-aba usando openpyxl.

Abas:
    1. Resumo Executivo   — KPIs principais
    2. Faturamento Mensal — evolucao no tempo
    3. Ranking Filiais    — top filiais com % do total
    4. Ranking Produtos   — top 20 produtos
    5. Top Clientes       — top clientes por LTV
    6. Faixa Etaria       — segmentacao demografica
    7. Categorias         — faturamento e margem por categoria
    8. Formas Pagamento   — distribuicao de forma de pagamento
    9. Insights           — texto com achados principais
"""

from pathlib import Path
from typing import Dict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule


# Estilos reutilizaveis
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="00695C")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="757575")
FONT_KPI_LABEL = Font(name="Calibri", size=10, color="757575")
FONT_KPI_VALUE = Font(name="Calibri", size=14, bold=True, color="00695C")

FILL_HEADER = PatternFill(start_color="00695C", end_color="00695C", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")
FILL_KPI = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _escrever_titulo_aba(ws, titulo: str, subtitulo: str = ""):
    """Escreve linha 1 com titulo e linha 2 com subtitulo em cada aba."""
    ws.cell(row=1, column=1, value=titulo).font = FONT_TITLE
    if subtitulo:
        ws.cell(row=2, column=1, value=subtitulo).font = FONT_SUBTITLE


def _escrever_dataframe(ws, df: pd.DataFrame, linha_inicio: int = 4,
                         cols_moeda: list = None, cols_percentual: list = None):
    """Escreve um DataFrame numa worksheet a partir de linha_inicio.

    - Cabecalho estilizado
    - Bordas nas celulas
    - Formatacao de moeda e percentual conforme parametros
    """
    cols_moeda = cols_moeda or []
    cols_percentual = cols_percentual or []

    # Cabecalho
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=linha_inicio, column=col_idx, value=str(col_name))
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN

    # Dados
    for row_idx, (_, row) in enumerate(df.iterrows(), start=linha_inicio + 1):
        for col_idx, col_name in enumerate(df.columns, start=1):
            valor = row[col_name]

            # Converte tipos pandas para nativos
            if pd.isna(valor):
                valor = None
            elif hasattr(valor, "item"):
                valor = valor.item()

            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = BORDER_THIN

            # Formatacao
            if col_name in cols_moeda:
                cell.number_format = 'R$ #,##0.00'
                cell.alignment = ALIGN_RIGHT
            elif col_name in cols_percentual:
                cell.number_format = "0.00\"%\""
                cell.alignment = ALIGN_RIGHT
            elif isinstance(valor, (int, float)):
                cell.number_format = "#,##0"
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT

    # Auto-fit aproximado
    for col_idx, col_name in enumerate(df.columns, start=1):
        letra = get_column_letter(col_idx)
        try:
            max_len = df[col_name].astype(str).map(len).max()
        except Exception:
            max_len = 10
        if pd.isna(max_len):
            max_len = 10
        ws.column_dimensions[letra].width = max(len(str(col_name)) + 2, int(max_len) + 3, 12)

    return linha_inicio + len(df) + 1


def gerar_relatorio_excel(
    kpis: Dict[str, float],
    faturamento_mensal: pd.DataFrame,
    faturamento_filial: pd.DataFrame,
    ranking_produtos: pd.DataFrame,
    top_clientes: pd.DataFrame,
    segmentacao_etaria: pd.DataFrame,
    faturamento_categoria: pd.DataFrame,
    faturamento_pagamento: pd.DataFrame,
    insights: list,
    caminho_saida: str,
):
    """Gera Excel completo do relatorio."""
    Path(caminho_saida).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # =============================================================
    # Aba 1: Resumo Executivo
    # =============================================================
    ws = wb.active
    ws.title = "Resumo Executivo"
    _escrever_titulo_aba(ws, "Resumo Executivo — PharmaMinas",
                          "Dados agregados de todo o periodo | Gerado automaticamente")

    kpis_lista = [
        ("Faturamento Total", kpis["faturamento_total"], "R$ #,##0.00"),
        ("Total de Vendas Concluidas", kpis["total_vendas"], "#,##0"),
        ("Total de Itens Vendidos", kpis["total_itens_vendidos"], "#,##0"),
        ("Ticket Medio", kpis["ticket_medio"], "R$ #,##0.00"),
        ("Margem Bruta Total", kpis["margem_bruta_total"], "R$ #,##0.00"),
        ("Percentual de Margem", kpis["percentual_margem"], "0.00\"%\""),
        ("Clientes Distintos", kpis["clientes_ativos"], "#,##0"),
        ("Produtos Distintos Vendidos", kpis["produtos_vendidos_distintos"], "#,##0"),
        ("Filiais Ativas", kpis["filiais_ativas"], "#,##0"),
    ]

    for i, (label, valor, formato) in enumerate(kpis_lista, start=4):
        cell_label = ws.cell(row=i, column=1, value=label)
        cell_label.font = FONT_KPI_LABEL
        cell_label.fill = FILL_KPI
        cell_label.alignment = ALIGN_LEFT
        cell_label.border = BORDER_THIN

        cell_val = ws.cell(row=i, column=2, value=valor)
        cell_val.font = FONT_KPI_VALUE
        cell_val.number_format = formato
        cell_val.alignment = ALIGN_RIGHT
        cell_val.border = BORDER_THIN
        cell_val.fill = FILL_KPI

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    # =============================================================
    # Aba 2: Faturamento Mensal
    # =============================================================
    ws = wb.create_sheet("Faturamento Mensal")
    _escrever_titulo_aba(ws, "Faturamento Mensal",
                          "Evolucao do faturamento agregado por mes")
    _escrever_dataframe(
        ws, faturamento_mensal,
        cols_moeda=["faturamento", "ticket_medio"],
    )

    # =============================================================
    # Aba 3: Ranking Filiais
    # =============================================================
    ws = wb.create_sheet("Ranking Filiais")
    _escrever_titulo_aba(ws, "Ranking de Filiais por Faturamento",
                          "Ordenado do maior para o menor")
    _escrever_dataframe(
        ws, faturamento_filial,
        cols_moeda=["faturamento", "ticket_medio"],
        cols_percentual=["percentual_do_total"],
    )

    # Escala de cor no faturamento
    if len(faturamento_filial) > 0:
        col_letter = get_column_letter(faturamento_filial.columns.get_loc("faturamento") + 1)
        rng = f"{col_letter}5:{col_letter}{4 + len(faturamento_filial)}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="min", start_color="B2DFDB",
            end_type="max", end_color="00695C",
        ))

    # =============================================================
    # Aba 4: Ranking Produtos
    # =============================================================
    ws = wb.create_sheet("Ranking Produtos")
    _escrever_titulo_aba(ws, "Ranking de Produtos",
                          "Top produtos por faturamento")
    _escrever_dataframe(
        ws, ranking_produtos,
        cols_moeda=["faturamento", "margem", "ticket_medio_item"],
        cols_percentual=["margem_percentual"],
    )

    # =============================================================
    # Aba 5: Top Clientes
    # =============================================================
    ws = wb.create_sheet("Top Clientes")
    _escrever_titulo_aba(ws, "Top Clientes por LTV",
                          "Lifetime Value ordenado do maior para o menor")

    top_clientes_export = top_clientes.copy()
    for col in ["primeira_compra", "ultima_compra"]:
        if col in top_clientes_export.columns:
            top_clientes_export[col] = top_clientes_export[col].dt.strftime("%Y-%m-%d")

    _escrever_dataframe(
        ws, top_clientes_export,
        cols_moeda=["ltv", "ticket_medio"],
    )

    # =============================================================
    # Aba 6: Faixa Etaria
    # =============================================================
    ws = wb.create_sheet("Faixa Etaria")
    _escrever_titulo_aba(ws, "Segmentacao por Faixa Etaria",
                          "Distribuicao do faturamento por idade dos clientes")
    _escrever_dataframe(
        ws, segmentacao_etaria,
        cols_moeda=["faturamento", "ticket_medio"],
        cols_percentual=["percentual"],
    )

    # =============================================================
    # Aba 7: Categorias
    # =============================================================
    ws = wb.create_sheet("Categorias")
    _escrever_titulo_aba(ws, "Faturamento por Categoria",
                          "Volume e margem por categoria de produto")
    _escrever_dataframe(
        ws, faturamento_categoria,
        cols_moeda=["faturamento", "margem"],
        cols_percentual=["margem_percentual", "percentual_do_total"],
    )

    # =============================================================
    # Aba 8: Formas de Pagamento
    # =============================================================
    ws = wb.create_sheet("Formas Pagamento")
    _escrever_titulo_aba(ws, "Distribuicao de Formas de Pagamento",
                          "Ordenado por faturamento")
    _escrever_dataframe(
        ws, faturamento_pagamento,
        cols_moeda=["faturamento"],
        cols_percentual=["percentual_do_total"],
    )

    # =============================================================
    # Aba 9: Insights
    # =============================================================
    ws = wb.create_sheet("Insights")
    _escrever_titulo_aba(ws, "Insights de Negocio",
                          "Achados principais da analise exploratoria")

    for i, insight in enumerate(insights, start=4):
        cell = ws.cell(row=i, column=1, value=f"{i - 3}. {insight}")
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.font = Font(name="Calibri", size=11, color="424242")
        ws.row_dimensions[i].height = 45

    ws.column_dimensions["A"].width = 110

    # Salvar
    wb.save(caminho_saida)
    return caminho_saida
